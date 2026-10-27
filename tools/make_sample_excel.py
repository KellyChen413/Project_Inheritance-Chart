"""以虛構範例產生戶役政二親等資料的 Excel 測試檔（模擬實際匯出可能出現的各種情形）
執行：python tools/make_sample_excel.py
- 範例1_文字.xlsx：所有儲存格皆為文字（含標題列）
- 範例1_數字日期.xlsx：日期被 Excel 轉成數字（0370512 → 370512）
- 範例1_日期格式.xlsx：日期為 Excel 日期格式儲存格
- 範例1_舊版.xls：Excel 97-2003（xlwt 產生）
- 範例2_LibreOffice.xls：由 LibreOffice 轉出的 .xls（含共用字串表、CONTINUE 等）
- 範例1_網頁格式.xls：部分系統匯出的「HTML 表格存成 .xls」
"""
import csv, datetime, subprocess, shutil
from pathlib import Path
import openpyxl, xlwt

root = Path(__file__).resolve().parent.parent
out = root / "data/excel"; out.mkdir(exist_ok=True)
DATE_COLS = {"出生日期", "死亡日期", "結婚日期"}
TITLE = "二親等關聯資料查詢結果（虛構範例）"

def rows_of(name):
    with open(root / "data" / name, encoding="utf-8-sig") as f:
        return list(csv.reader(f))

def roc_to_date(v):
    y, m, d = int(v[:-4]), int(v[-4:-2]), int(v[-2:])
    return datetime.date(y + 1911, m, d)

def xlsx(rows, path, mode):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "查詢結果"
    ws.append([TITLE]); ws.append([])
    hdr = rows[0]; ws.append(hdr)
    for r in rows[1:]:
        cells = []
        for h, v in zip(hdr, r):
            if h in DATE_COLS and v and mode == "number": cells.append(int(v))
            elif h in DATE_COLS and v and mode == "date": cells.append(roc_to_date(v))
            elif h == "項目" and mode != "text": cells.append(int(v))
            else: cells.append(v)
        ws.append(cells)
    if mode == "date":
        for row in ws.iter_rows(min_row=4):
            for c in row:
                if isinstance(c.value, datetime.date): c.number_format = "yyyy/mm/dd"
    wb.save(path)

def xls(rows, path):
    wb = xlwt.Workbook(encoding="utf-8"); ws = wb.add_sheet("查詢結果")
    ws.write(0, 0, TITLE)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            if rows[0][j] in DATE_COLS and v and i: ws.write(i + 2, j, int(v))   # 數字型民國日期
            else: ws.write(i + 2, j, v)
    wb.save(str(path))

def html_xls(rows, path):
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;")
    body = "".join("<tr>" + "".join(f"<td>{esc(v)}</td>" for v in r) + "</tr>" for r in rows)
    path.write_text(f'<html><head><meta charset="utf-8"></head><body><table><tr><td colspan="16">{TITLE}</td></tr>{body}</table></body></html>', encoding="utf-8")

r1, r2 = rows_of("範例1_王大明_二親等資料.csv"), rows_of("範例2_王淑娟_二親等資料.csv")
xlsx(r1, out / "範例1_文字.xlsx", "text")
xlsx(r1, out / "範例1_數字日期.xlsx", "number")
xlsx(r1, out / "範例1_日期格式.xlsx", "date")
xls(r1, out / "範例1_舊版.xls")
html_xls(r1, out / "範例1_網頁格式.xls")
# LibreOffice 轉檔
tmp = out / "_tmp"; tmp.mkdir(exist_ok=True)
xlsx(r2, tmp / "範例2_LibreOffice.xlsx", "text")
subprocess.run(["soffice", "--headless", "--convert-to", "xls", "--outdir", str(out), str(tmp / "範例2_LibreOffice.xlsx")], check=True, capture_output=True)
shutil.rmtree(tmp)
print(sorted(p.name for p in out.iterdir()))
